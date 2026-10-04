# -*- coding: utf-8 -*-
"""
================================================================================
 자가용 출장비 자동 계산기 + 증빙 보고서(PDF·엑셀)  v5
================================================================================
 주소만 입력하면 ① 카카오내비 길찾기 API로 경로 후보(거리·시간·통행료)를 조회하고,
 ② 오피넷 실시간 유가를 반영해 ③ 사내 여비 규정(시간단위 포함)에 따라
 운임·식비·일비를 산출한 뒤, ④ 실제 주행 경로 지도가 포함된 증빙 보고서 PDF와
 ⑤ 사내 양식(국내출장 여비 집행 리스트) 엑셀을 생성합니다.

 [비용 없는 구성]  카카오맵(Local/지오코딩)은 사용하지 않습니다.
   · 주소→좌표 : OSM Nominatim(무료, 키 불필요) 기본
                 + VWorld(무료, 일 40,000건) / 도로명주소 API(무료) 키를 넣으면 정확도 향상
   · 경로/거리/통행료 : 카카오내비 길찾기 API(보유 키) → 실패 시 OSM(OSRM) 무료 폴백
   · 실시간 유가 : 오피넷(무료)

 [근거 규정] 여비업무 처리 매뉴얼(2025.09.23.)
   3.운임 나) 자가용  : 연료비 = 거리(km) ÷ 연비 × 유가, 통행료·주차료 실비(주차료 1일 25,000원)
   2.여비지급기준     : 근무지내 4시간 미만 10,000원 / 4시간 이상 20,000원
   4.식비             : 4시간 미만 1/3 · 4~6시간 2/3 · 6시간 이상 전액 (1일 25,000원)
   5.일비             : 1일 25,000원 정액
   연비·전비          : 휘발유 11.97 | 경유 12.52 | LPG 8.83 | PHEV 15.37 | HEV 10.61
                        (km/L), 전기 2.84 (km/kWh), 수소 94.9 (km/kg)

 [실행]  pip install -r requirements.txt  →  streamlit run app.py
================================================================================
"""

from __future__ import annotations

import datetime as dt
import inspect
import io
import json
import math
import os
import urllib.error
import urllib.parse
import urllib.request
from dataclasses import dataclass
from pathlib import Path
from typing import Any, Optional

import pandas as pd
import streamlit as st

BASE_DIR = Path(__file__).resolve().parent
ASSETS_DIR = BASE_DIR / "assets"
# GitHub 웹 업로드 시 assets/ 폴더 없이 루트에 올라가도 동작하도록 폴백
ASSET_DIRS = [ASSETS_DIR, BASE_DIR]


def _asset_path(name: str):
    """assets/ 우선, 없으면 저장소 루트에서 파일을 찾는다."""
    for d in ASSET_DIRS:
        p = d / name
        if p.exists():
            return p
    return None

# ==============================================================================
# 1. 규정 상수
# ==============================================================================
PARKING_DAILY_CAP: int = 25_000
WITHIN_WORK_UNDER_4H: int = 10_000
WITHIN_WORK_OVER_4H: int = 20_000
MEAL_ALLOWANCE_PER_DAY: int = 25_000
DAILY_ALLOWANCE_PER_DAY: int = 25_000

#: 오피넷 유가정보 API
OPINET_AVG_ALL_URL = "https://www.opinet.co.kr/api/avgAllPrice.do"
OPINET_AVG_SIDO_URL = "https://www.opinet.co.kr/api/avgSidoPrice.do"

#: 카카오내비 길찾기 (카카오맵 Local API와 별개 서비스 — 지오코딩에는 사용하지 않음)
KAKAO_DIRECTIONS_URL = "https://apis-navi.kakaomobility.com/v1/directions"

#: 무료 지오코딩
VWORLD_GEOCODE_URL = "https://api.vworld.kr/req/address"
JUSO_ADDR_URL = "https://business.juso.go.kr/addrlink/addrLinkApi.do"
OSM_NOMINATIM_URL = "https://nominatim.openstreetmap.org/search"
OSM_OSRM_URL = "http://router.project-osrm.org/route/v1/driving"
USER_AGENT = "trip-expense-calculator/5.0 (internal expense report tool)"

OPINET_PROD_CODE = {"보통휘발유": "B027", "고급휘발유": "B034", "자동차경유": "D047",
                    "실내등유": "C004", "자동차부탄": "K015"}
ROUTE_PRIORITIES = {"추천 경로": "RECOMMEND", "최단 시간": "TIME",
                    "최단 거리": "DISTANCE", "큰길 우선": "MAIN_ROAD"}
ROUNDING_MODES = ("원 단위 절사(버림)", "원 단위 반올림", "원 단위 올림")
TRIP_SCOPES = ("근무지외 국내출장", "근무지내 국내출장")


@dataclass(frozen=True)
class FuelSpec:
    name: str
    efficiency: float
    unit: str
    price_unit: str
    opinet_prodcd: Optional[str]
    price_source: str
    fallback_price: float = 0.0
    note: str = ""


FUEL_SPECS: dict[str, FuelSpec] = {
    "휘발유": FuelSpec("휘발유", 11.97, "km/L", "원/L", OPINET_PROD_CODE["보통휘발유"],
                     "오피넷 보통휘발유 평균가"),
    "고급휘발유": FuelSpec("고급휘발유", 11.97, "km/L", "원/L", OPINET_PROD_CODE["고급휘발유"],
                        "오피넷 고급휘발유 평균가", note="연비표상 휘발유와 동일 기준"),
    "경유": FuelSpec("경유", 12.52, "km/L", "원/L", OPINET_PROD_CODE["자동차경유"],
                    "오피넷 자동차용경유 평균가"),
    "LPG(부탄)": FuelSpec("LPG(부탄)", 8.83, "km/L", "원/L", OPINET_PROD_CODE["자동차부탄"],
                       "오피넷 자동차부탄 평균가"),
    "하이브리드": FuelSpec("하이브리드", 10.61, "km/L", "원/L", OPINET_PROD_CODE["보통휘발유"],
                       "오피넷 보통휘발유 평균가"),
    "플러그인 하이브리드": FuelSpec("플러그인 하이브리드", 15.37, "km/L", "원/L",
                          OPINET_PROD_CODE["보통휘발유"], "오피넷 보통휘발유 평균가",
                          note="전기 주행분 포함 복합 연비"),
    "전기": FuelSpec("전기", 2.84, "km/kWh", "원/kWh", None,
                    "환경부 무공해차 통합누리집 충전요금", 350.0,
                    note="오피넷 미제공 — 충전요금 직접 입력"),
    "수소": FuelSpec("수소", 94.9, "km/kg", "원/kg", None, "수소충전소 고시가", 9500.0,
                    note="오피넷 미제공 — 충전요금 직접 입력"),
}


# ==============================================================================
# 2. 계산 로직
# ==============================================================================

def round_won(value: float, mode: str = ROUNDING_MODES[0]) -> int:
    if value is None or (isinstance(value, float) and math.isnan(value)):
        return 0
    if mode == "원 단위 반올림":
        return int(round(value))
    if mode == "원 단위 올림":
        return int(math.ceil(value))
    return int(math.floor(value))


def calc_fuel_cost(distance_km: float, efficiency: float, unit_price: float,
                   rounding_mode: str = ROUNDING_MODES[0]) -> int:
    """연료비 = 출장거리(km) ÷ 연비 × 유가"""
    if efficiency <= 0:
        raise ValueError("연비는 0보다 커야 합니다.")
    return round_won((float(distance_km) / float(efficiency)) * float(unit_price), rounding_mode)


def apply_parking_cap(parking_fee: float, days: int,
                      daily_cap: int = PARKING_DAILY_CAP) -> tuple[int, bool, int]:
    days = max(int(days), 1)
    cap_total = daily_cap * days
    fee = int(max(float(parking_fee or 0), 0))
    return (cap_total, True, cap_total) if fee > cap_total else (fee, False, cap_total)


def meal_ratio(hours: float) -> tuple[int, int, str]:
    """식비 지급 비율 (매뉴얼 4.식비 나)"""
    if hours < 4:
        return 1, 3, "출장시간 4시간 미만 → 식비 1/3"
    if hours < 6:
        return 2, 3, "출장시간 4시간 이상~6시간 미만 → 식비 2/3"
    return 1, 1, "출장시간 6시간 이상 → 식비 전액"


def calc_meal_allowance(hours: float, days: int = 1, per_day: int = MEAL_ALLOWANCE_PER_DAY,
                        rounding_mode: str = ROUNDING_MODES[0]) -> tuple[int, str]:
    n, d, note = meal_ratio(hours)
    return round_won(per_day * max(int(days), 1) * n / d, rounding_mode), note


def calc_daily_allowance(days: int = 1, per_day: int = DAILY_ALLOWANCE_PER_DAY) -> int:
    return per_day * max(int(days), 1)


def calc_within_work_allowance(hours: float) -> tuple[int, str]:
    """근무지내 국내출장 출장비 (매뉴얼 2.여비지급기준 가)"""
    if hours >= 4:
        return WITHIN_WORK_OVER_4H, "출장시간 4시간 이상 → 20,000원"
    return WITHIN_WORK_UNDER_4H, "출장시간 4시간 미만 → 10,000원"


def hours_to_text(hours: float) -> str:
    h = int(hours)
    m = int(round((hours - h) * 60))
    if m == 60:
        h, m = h + 1, 0
    return f"{h}시간 {m}분" if m else f"{h}시간"


def build_route_text(origin: str, waypoints: list[str], destination: str) -> str:
    parts = [origin.strip() or "(미입력)"]
    parts += [f"경유: {w.strip()}" for w in waypoints if w and w.strip()]
    parts.append(destination.strip() or "(미입력)")
    return " → ".join(parts)


# ==============================================================================
# 3. 인증키 / HTTP
# ==============================================================================

def get_secret(name: str, default: str = "") -> str:
    try:
        if name in st.secrets:
            return str(st.secrets[name]).strip()
    except Exception:
        pass
    return str(os.environ.get(name, default)).strip()


def _http_text(url: str, headers: Optional[dict] = None, timeout: int = 25) -> str:
    req = urllib.request.Request(url, headers=headers or {"User-Agent": USER_AGENT})
    with urllib.request.urlopen(req, timeout=timeout) as resp:
        return resp.read().decode("utf-8", "replace")


# ==============================================================================
# 4. 오피넷 실시간 유가
# ==============================================================================

@st.cache_data(ttl=1800, show_spinner=False)
def fetch_opinet_all(cert_key: str) -> dict:
    if not cert_key:
        return {}
    try:
        oils = json.loads(_http_text(
            f"{OPINET_AVG_ALL_URL}?out=json&certkey={urllib.parse.quote(cert_key)}")
        ).get("RESULT", {}).get("OIL", [])
        if isinstance(oils, dict):
            oils = [oils]
        return {o.get("PRODCD"): {"price": float(o.get("PRICE", 0)), "name": o.get("PRODNM", ""),
                                 "date": o.get("TRADE_DT", ""), "diff": o.get("DIFF", "")}
                for o in oils}
    except Exception:
        return {}


@st.cache_data(ttl=1800, show_spinner=False)
def fetch_opinet_by_sido(cert_key: str) -> dict:
    if not cert_key:
        return {}
    try:
        oils = json.loads(_http_text(
            f"{OPINET_AVG_SIDO_URL}?out=json&certkey={urllib.parse.quote(cert_key)}")
        ).get("RESULT", {}).get("OIL", [])
        if isinstance(oils, dict):
            oils = [oils]
        out: dict[str, dict[str, float]] = {}
        for o in oils:
            out.setdefault(o.get("SIDONM", ""), {})[o.get("PRODCD", "")] = float(o.get("PRICE", 0))
        return out
    except Exception:
        return {}


def opinet_price(cert_key: str, prodcd: Optional[str], sido: str = "") -> Optional[float]:
    if not cert_key or not prodcd:
        return None
    if sido:
        v = fetch_opinet_by_sido(cert_key).get(sido, {}).get(prodcd)
        if v:
            return float(v)
    d = fetch_opinet_all(cert_key)
    return float(d[prodcd]["price"]) if prodcd in d else None


# ==============================================================================
# 5. 무료 주소 → 좌표 (카카오맵 Local API 미사용)
# ==============================================================================

def _address_variants(address: str) -> list[str]:
    a = " ".join(address.split())
    out = [a]
    for token in ("경상북도", "경북", "서울특별시", "서울", "부산광역시", "부산", "대전광역시",
                  "대전", "대구광역시", "대구", "인천광역시", "인천", "광주광역시", "광주",
                  "울산광역시", "울산", "세종특별자치시", "세종", "경기도", "강원특별자치도",
                  "강원도", "충청북도", "충북", "충청남도", "충남", "전라북도", "전북",
                  "전라남도", "전남", "경상남도", "경남", "제주특별자치도", "제주"):
        if a.startswith(token + " "):
            out.append(a[len(token) + 1:])
            break
    parts = a.split()
    if len(parts) >= 2:
        out.append(" ".join(parts[:-1]))
    if len(parts) >= 3:
        out.append(" ".join(parts[:2]))
    seen, uniq = set(), []
    for v in out:
        if v and v not in seen:
            seen.add(v)
            uniq.append(v)
    return uniq


#: VWorld 마지막 오류 메시지(화면 로그 표시용)
_VWORLD_LAST_ERROR: str = ""


def vworld_geocode(address: str, api_key: str, domain: str = "") -> Optional[dict]:
    """VWorld **Geocoder API 2.0** (무료, 일 40,000건).

    ▸ 인증키 발급: https://www.vworld.kr → 오픈API → 인증키 발급/관리
      · 신청할 API: **Geocoder API** (버전 2.0)  ← 주소→좌표 변환
      · 인증키 종류: 개발키(6개월, 즉시 발급) / 운영키(기관·사업자, 심의)
      · **서비스 URL(도메인) 등록 필수** → 등록한 도메인을 domain 파라미터로 전달
        (예: 로컬 테스트면 http://localhost:8501 → domain=localhost:8501)
    ▸ 요청: GET https://api.vworld.kr/req/address
        ?service=address&request=getcoord&version=2.0&crs=EPSG:4326
        &type=road|parcel&address=주소&format=json&errorformat=json&key=인증키&domain=등록도메인
    """
    global _VWORLD_LAST_ERROR
    if not api_key:
        return None
    # 도메인 후보: 등록값 그대로 → 끝 슬래시 제거 → 스킴 제거 → 생략(빈값) 순으로 재시도
    domains: list = []
    if domain:
        _d = domain.strip()
        for _c in (_d, _d.rstrip("/"), _d.split("://", 1)[-1].rstrip("/")):
            if _c and _c not in domains:
                domains.append(_c)
    domains.append("")
    for addr_type in ("road", "parcel"):
        for dm in domains:
            try:
                params = {"service": "address", "request": "getcoord", "version": "2.0",
                          "crs": "EPSG:4326", "type": addr_type, "address": address,
                          "format": "json", "errorformat": "json", "key": api_key}
                if dm:
                    params["domain"] = dm
                resp = json.loads(_http_text(
                    f"{VWORLD_GEOCODE_URL}?{urllib.parse.urlencode(params)}",
                    timeout=10)).get("response", {})
                status = resp.get("status")
                if status == "OK":
                    pt = resp["result"]["point"]
                    _VWORLD_LAST_ERROR = ""
                    return {"lng": float(pt["x"]), "lat": float(pt["y"]),
                            "label": (resp.get("refined") or {}).get("text", address),
                            "provider": "VWorld"}
                err = resp.get("error") or {}
                code = err.get("code", status or "?")
                text = err.get("text", "")
                if code == "INCORRECT_KEY":
                    _VWORLD_LAST_ERROR = (
                        f"VWorld 인증키 도메인 불일치({code}). 발급 시 등록한 서비스 URL을 "
                        f"VWORLD_DOMAIN 에 그대로 입력하세요. (현재 domain={dm!r})")
                else:
                    _VWORLD_LAST_ERROR = f"VWorld {code} {text[:80]}"
            except Exception as exc:  # noqa: BLE001
                _code = getattr(exc, "code", None)
                try:
                    _body = exc.read().decode("utf-8", "replace").strip()[:220]
                except Exception:
                    _body = ""
                _VWORLD_LAST_ERROR = (
                    f"VWorld 호출 실패(HTTP {_code or type(exc).__name__})"
                    + (f" 응답: {_body}" if _body else "")
                    + f" | domain={dm!r}, type={addr_type}")
                continue
    return None


def juso_geocode(address: str, api_key: str) -> Optional[dict]:
    """행안부 도로명주소 검색 API (무료). business.juso.go.kr 에서 승인키 발급."""
    if not api_key:
        return None
    try:
        q = urllib.parse.urlencode({"confmKey": api_key, "currentPage": 1,
                                    "countPerPage": 1, "keyword": address,
                                    "resultType": "json"})
        juso = json.loads(_http_text(f"{JUSO_ADDR_URL}?{q}")).get("results", {}).get("juso", [])
        if juso:
            j = juso[0]
            if j.get("entX") and j.get("entY"):
                return {"lng": float(j["entX"]), "lat": float(j["entY"]),
                        "label": j.get("roadAddr") or j.get("jibunAddr") or address,
                        "provider": "도로명주소 API"}
    except Exception:
        return None
    return None


def osm_geocode(address: str) -> Optional[dict]:
    """OpenStreetMap Nominatim (무료, 키 불필요) — 주소 변형 다중 시도."""
    for q in _address_variants(address):
        try:
            url = f"{OSM_NOMINATIM_URL}?" + urllib.parse.urlencode(
                {"q": q, "format": "json", "limit": 1, "countrycodes": "kr",
                 "addressdetails": 1})
            arr = json.loads(_http_text(url))
            if arr:
                return {"lng": float(arr[0]["lon"]), "lat": float(arr[0]["lat"]),
                        "label": arr[0].get("display_name", q), "provider": "OSM(Nominatim)"}
        except Exception:
            continue
    return None


def geocode(address: str, vworld_key: str = "", juso_key: str = "",
            vworld_domain: str = "") -> Optional[dict]:
    """무료 지오코딩 체인: VWorld(Geocoder 2.0) → 도로명주소 → OSM(Nominatim)."""
    return (vworld_geocode(address, vworld_key, vworld_domain)
            or juso_geocode(address, juso_key)
            or osm_geocode(address))


# ==============================================================================
# 6. 경로 (카카오내비 → OSRM 무료 폴백)
# ==============================================================================

def kakao_routes(points: list[dict], rest_key: str, priority: str = "RECOMMEND",
                 alternatives: bool = True) -> list[dict]:
    if not rest_key or len(points) < 2:
        return []
    try:
        params = {"origin": f"{points[0]['lng']},{points[0]['lat']}",
                  "destination": f"{points[-1]['lng']},{points[-1]['lat']}",
                  "priority": priority, "car_fuel": "GASOLINE",
                  "summary": "false", "road_details": "true"}
        if alternatives:
            params["alternatives"] = "true"
        mids = points[1:-1][:5]
        if mids:
            params["waypoints"] = "|".join(f"{p['lng']},{p['lat']}" for p in mids)
        raw = _http_text(f"{KAKAO_DIRECTIONS_URL}?{urllib.parse.urlencode(params)}",
                         {"Authorization": f"KakaoAK {rest_key}", "User-Agent": USER_AGENT})
        out: list[dict] = []
        for rt in json.loads(raw).get("routes", []):
            if rt.get("result_code") != 0:
                continue
            s = rt["summary"]
            poly, names = [], []
            for sec in rt.get("sections", []):
                for rd in sec.get("roads", []):
                    v = rd.get("vertexes") or []
                    for i in range(0, len(v) - 1, 2):
                        poly.append((float(v[i]), float(v[i + 1])))
                    if rd.get("name"):
                        names.append(rd["name"])
            if not poly:
                poly = [(p["lng"], p["lat"]) for p in points]
            out.append({"distance_km": s["distance"] / 1000.0,
                        "duration_min": s["duration"] / 60.0,
                        "toll_fee": int((s.get("fare") or {}).get("toll", 0)),
                        "polyline": poly, "road_names": names,
                        "priority": s.get("priority", priority),
                        "provider": "카카오내비 길찾기"})
        out.sort(key=lambda r: r["distance_km"])
        uniq, seen = [], set()
        for r in out:
            k = (round(r["distance_km"], 1), r["toll_fee"])
            if k not in seen:
                seen.add(k)
                uniq.append(r)
        return uniq[:5]
    except Exception:
        return []


def osrm_route(points: list[dict]) -> Optional[dict]:
    if len(points) < 2:
        return None
    try:
        coords = ";".join(f"{p['lng']},{p['lat']}" for p in points)
        r = json.loads(_http_text(
            f"{OSM_OSRM_URL}/{coords}?overview=full&geometries=geojson"
            f"&alternatives=false&steps=false"))
        if r.get("code") != "Ok":
            return None
        rt = r["routes"][0]
        return {"distance_km": rt["distance"] / 1000.0, "duration_min": rt["duration"] / 60.0,
                "toll_fee": 0,
                "polyline": [(float(x), float(y)) for x, y in rt["geometry"]["coordinates"]],
                "road_names": [], "priority": "OSRM", "provider": "OSM(OSRM) 무료 폴백"}
    except Exception:
        return None


def resolve_routes(origin: str, waypoints: list[str], destination: str,
                   kakao_key: str, vworld_key: str = "", juso_key: str = "",
                   priority: str = "RECOMMEND",
                   vworld_domain: str = "") -> tuple[list[dict], list[str]]:
    logs: list[str] = []
    addrs = [origin] + [w for w in waypoints if w.strip()] + [destination]
    points: list[dict] = []
    for a in addrs:
        p = geocode(a, vworld_key, juso_key, vworld_domain)
        if not p:
            logs.append(f"❌ 주소를 찾지 못했습니다: {a}")
            return [], logs
        logs.append(f"📍 {a} → {p['label'][:70]} ({p['provider']})")
        points.append(p)
    if _VWORLD_LAST_ERROR and vworld_key:
        logs.append(f"⚠️ {_VWORLD_LAST_ERROR} → OSM으로 대체했습니다.")

    routes = kakao_routes(points, kakao_key, priority=priority, alternatives=True)
    if routes:
        logs.append(f"🛣️ 카카오내비 길찾기 — 경로 후보 {len(routes)}개 조회")
    else:
        single = osrm_route(points)
        if single:
            routes = [single]
            logs.append("🛣️ 카카오 길찾기 실패/미설정 → OSM(OSRM) 무료 폴백 (통행료 직접 입력)")
    if not routes:
        logs.append("❌ 경로 계산 실패 — 총 주행거리를 직접 입력해 주세요.")
        return [], logs
    for r in routes:
        r["points"] = points
    return routes, logs


# ==============================================================================
# 7. 경로 지도 렌더링
# ==============================================================================

@st.cache_data(ttl=86400, show_spinner=False)
def _load_basemap() -> Optional[dict]:
    for name in ("skorea_provinces_simple.json", "KOR.geo.json"):
        p = _asset_path(name)
        if p is not None:
            try:
                return json.loads(p.read_text(encoding="utf-8"))
            except Exception:
                continue
    return None


def _rings(geometry: dict) -> list[list[tuple[float, float]]]:
    out: list[list[tuple[float, float]]] = []
    coords = geometry.get("coordinates") or []
    gtype = geometry.get("type")
    polys = [coords] if gtype == "Polygon" else (coords if gtype == "MultiPolygon" else [])
    for poly in polys:
        for ring in poly:
            out.append([(float(x), float(y)) for x, y in ring])
    return out


def render_route_map(route: dict, labels: list[str], title: str = "출장 경로",
                     routes_all: Optional[list[dict]] = None) -> Optional[bytes]:
    try:
        import matplotlib
        matplotlib.use("Agg")
        from matplotlib import font_manager
        from matplotlib.figure import Figure
        from matplotlib.patches import Polygon as MplPolygon
    except Exception:
        return None

    poly = route.get("polyline") or []
    pts = route.get("points") or []
    if not poly and not pts:
        return None

    for fp in (ASSETS_DIR / "NanumGothic.ttf",
               BASE_DIR / "NanumGothic.ttf",
               Path("/usr/share/fonts/truetype/nanum/NanumGothic.ttf")):
        if fp.exists():
            try:
                font_manager.fontManager.addfont(str(fp))
                matplotlib.rcParams["font.family"] = "NanumGothic"
            except Exception:
                pass
            break
    matplotlib.rcParams["axes.unicode_minus"] = False

    lons = [p[0] for p in poly] or [p["lng"] for p in pts]
    lats = [p[1] for p in poly] or [p["lat"] for p in pts]
    for p in pts:
        lons.append(p["lng"])
        lats.append(p["lat"])
    minx, maxx, miny, maxy = min(lons), max(lons), min(lats), max(lats)
    padx = max((maxx - minx) * 0.06, 0.02)
    pady = max((maxy - miny) * 0.13, 0.04)
    minx, maxx, miny, maxy = minx - padx, maxx + padx, miny - pady, maxy + pady

    cosf = max(math.cos(math.radians((miny + maxy) / 2.0)), 0.2)
    need = 1.55 * (maxy - miny) / cosf
    if (maxx - minx) < need:
        extra = (need - (maxx - minx)) / 2.0
        minx -= extra
        maxx += extra

    fig = Figure(figsize=(9.6, 6.4), dpi=150)
    ax = fig.add_subplot(111)
    fig.patch.set_facecolor("white")
    ax.set_facecolor("#F7FAFC")

    base = _load_basemap()
    if base:
        for feat in base.get("features", []):
            for ring in _rings(feat.get("geometry") or {}):
                ax.add_patch(MplPolygon(ring, closed=True, facecolor="#E7EEF4",
                                        edgecolor="#FFFFFF", linewidth=0.7, zorder=1))

    for other in (routes_all or []):
        if other is route:
            continue
        op = other.get("polyline") or []
        if len(op) > 1:
            ax.plot([q[0] for q in op], [q[1] for q in op], color="#9AA7B4",
                    linewidth=1.3, linestyle=(0, (4, 3)), zorder=2)

    if len(poly) > 1:
        xs, ys = [p[0] for p in poly], [p[1] for p in poly]
        ax.plot(xs, ys, color="white", linewidth=5.0, solid_capstyle="round",
                solid_joinstyle="round", zorder=3)
        ax.plot(xs, ys, color="#D64545", linewidth=2.3, solid_capstyle="round",
                solid_joinstyle="round", zorder=4)

    styles = [("#1E63B0", "출발"), ("#E08A1E", "경유"), ("#C0392B", "도착")]
    for i, p in enumerate(pts):
        c, tag = styles[0] if i == 0 else (styles[2] if i == len(pts) - 1 else styles[1])
        ax.plot(p["lng"], p["lat"], marker="o", markersize=9, color=c,
                markeredgecolor="white", markeredgewidth=1.8, zorder=6)
        name = labels[i] if i < len(labels) else tag
        if len(name) > 20:
            name = name[:19] + "…"
        frac = (p["lng"] - minx) / max(maxx - minx, 1e-9)
        off, ha = ((-10, 9), "right") if frac > 0.62 else ((10, 9), "left")
        ax.annotate(f"{tag}\n{name}", (p["lng"], p["lat"]), textcoords="offset points",
                    xytext=off, ha=ha, fontsize=8.0, color="#22303C", zorder=7,
                    bbox=dict(boxstyle="round,pad=0.32", fc="white", ec=c, lw=0.9, alpha=0.94))

    ax.set_xlim(minx, maxx)
    ax.set_ylim(miny, maxy)
    ax.set_aspect(1.0 / cosf, adjustable="box")
    ax.tick_params(labelsize=7, colors="#8A97A3")
    for sp in ax.spines.values():
        sp.set_color("#CBD5DE")
    ax.grid(color="#DCE5EC", linewidth=0.6, linestyle=":", zorder=0)
    ax.set_title(title, fontsize=13.5, color="#1F3B57", pad=11, fontweight="bold")

    info = (f"총 {route['distance_km']:,.1f} km   ·   약 {hours_to_text(route['duration_min'] / 60)}"
            f"   ·   통행료 {route.get('toll_fee', 0):,}원   ·   경로: {route.get('provider', '')}")
    ax.text(0.012, 0.022, info, transform=ax.transAxes, fontsize=8.6, color="#22303C",
            bbox=dict(boxstyle="round,pad=0.42", fc="#FFFFFF", ec="#B9C7D3", lw=0.8, alpha=0.95),
            zorder=8)

    buf = io.BytesIO()
    fig.subplots_adjust(left=0.055, right=0.985, top=0.925, bottom=0.085)
    fig.savefig(buf, format="png", dpi=150, facecolor="white")
    return buf.getvalue()


# ==============================================================================
# 8. 정산 텍스트
# ==============================================================================

def build_receipt_text(ctx: dict) -> str:
    if ctx["scope"] == "근무지내 국내출장":
        detail = (f"  - 출장비   : {ctx['within_work_cost']:,} 원 ({ctx['within_work_note']})\n"
                  f"              ※ 근무지내 출장은 운임·일비·식비·숙박비를 별도 지급하지 않습니다.")
        if ctx["toll_fee"]:
            detail += f"\n  - 통행료   : {ctx['toll_fee']:,} 원 (불가피한 유료도로 통행료)"
    else:
        detail = (
            f"  - 연 료 비 : {ctx['distance_km']:,.1f} km ÷ {ctx['efficiency']:,.2f} "
            f"{ctx['eff_unit']} × {ctx['unit_price']:,.2f} {ctx['price_unit']}"
            f" = {ctx['fuel_cost']:,} 원 ({ctx['rounding_mode']})\n"
            f"  - 통행료   : {ctx['toll_fee']:,} 원 (고속도로 통행영수증 실비)\n"
            f"  - 주차료   : {ctx['parking_fee']:,} 원"
            + (f" (1일 상한 25,000원 × {ctx['days']}일 적용)" if ctx["parking_capped"] else "")
            + f"\n  - 소계(운임): {ctx['transport_total']:,} 원")
        if ctx["include_allowances"]:
            detail += (f"\n  - 식 비    : {ctx['meal_cost']:,} 원 ({ctx['meal_note']})\n"
                       f"  - 일 비    : {ctx['daily_cost']:,} 원 (1일 25,000원 × {ctx['days']}일)")
    return f"""■ 자가용 이용 출장비 정산 내역

  - 출장자   : {ctx['traveler'] or '(미입력)'}  {ctx.get('dept','')}
  - 출장구분 : {ctx['scope']}
  - 출장기간 : {ctx['start_dt_text']} ~ {ctx['end_dt_text']} ({ctx['days']}일간)
  - 출장시간 : {ctx['travel_time_text']}
  - 출장목적 : {ctx['purpose'] or '-'}
  - 출장경로 : {ctx['route_text']}
  - 차종/유종: {ctx['fuel_name']} (기준 연비 {ctx['efficiency']:,.2f} {ctx['eff_unit']})
  - 주행거리 : {ctx['distance_km']:,.1f} km
  - 적용유가 : {ctx['unit_price']:,.2f} {ctx['price_unit']}  ※ {ctx['price_source']}

[산출 근거]
{detail}

[청구 총액]
  - 합계 : {ctx['total']:,} 원

[첨부 증빙서류]
  1. 고속도로 통행영수증 (통행료 발생 시)
  2. 출장지 소재 주유소에서 결제한 신용카드매출전표 (연료비)
  3. 주차영수증 (주차료 발생 시)
  ※ 자가용 동승자에게는 연료비·통행료·주차료를 지급하지 않습니다.
  ※ 2인 이상 동행 출장 시 1대 차량 이용이 원칙입니다.
"""


# ==============================================================================
# 9. 엑셀 (사내 양식: 국내출장 여비 집행 리스트)
# ==============================================================================

MONEY_FMT = '_-* #,##0_-;\\-* #,##0_-;_-* "-"_-;_-@_-'
EXCEL_COLS = ["사번", "성명", "부서", "직급", "기간시작", "기간끝", "사유(용무)", "행선지",
              "업무용차량", "출장여비", "일비", "식비", "교통비",
              "출장시간", "주행거리(km)", "연료비", "통행료", "주차료"]
COL_WIDTHS = {"A": 3.0, "B": 10.5, "C": 10.0, "D": 16.5, "E": 10.0, "F": 17.5, "G": 17.5,
              "H": 46.0, "I": 22.0, "J": 11.0, "K": 12.0, "L": 10.0, "M": 10.0, "N": 12.0,
              "O": 11.0, "P": 12.5, "Q": 12.0, "R": 12.0, "S": 12.0}


def build_excel(trips: list[dict], meta: dict) -> bytes:
    """사내 양식(국내출장 여비 집행 리스트) 엑셀 생성.

    trips: 정산 목록(각 항목은 아래 키 포함)
      emp_no, traveler, dept, rank, start_dt_text, end_dt_text, purpose, destination,
      company_car, total, daily, meal, transport, travel_time_text, distance_km,
      fuel_cost, toll_fee, parking_fee, scope, route_text, fuel_name, efficiency,
      unit_price, route_provider, bank, account, holder
    meta : {title, budget_label, budget_value}
    """
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
    from openpyxl.utils import get_column_letter

    thin = Side(style="thin", color="9E9E9E")
    border = Border(left=thin, right=thin, top=thin, bottom=thin)
    center = Alignment(horizontal="center", vertical="center", wrap_text=True)
    left = Alignment(horizontal="left", vertical="center", wrap_text=True)
    right = Alignment(horizontal="right", vertical="center")

    wb = Workbook()
    ws = wb.active
    ws.title = "출장비 집행 리스트"

    for col, w in COL_WIDTHS.items():
        ws.column_dimensions[col].width = w

    last_col = get_column_letter(len(EXCEL_COLS) + 1)          # B..S
    n = len(trips)
    header_row = 5
    first_data = 6
    last_data = first_data + n - 1
    total_row = last_data + 1

    # 제목 (B1:S3 병합)
    ws.merge_cells(f"B1:{last_col}3")
    ws["B1"] = meta.get("title", "국내출장 여비 집행 리스트")
    ws["B1"].font = Font(name="HY헤드라인M", size=26)
    ws["B1"].alignment = center
    for r in (1, 2, 3):
        ws.row_dimensions[r].height = 13.5

    # 예산과목
    ws["B4"] = meta.get("budget_label", "예산과목")
    ws["B4"].font = Font(name="맑은 고딕", size=11)
    ws.merge_cells(f"C4:{last_col}4")
    ws["C4"] = meta.get("budget_value", "")
    ws["C4"].font = Font(name="맑은 고딕", size=11)
    ws["C4"].alignment = left

    # 헤더
    for i, h in enumerate(EXCEL_COLS):
        c = ws.cell(header_row, i + 2, h)
        c.font = Font(name="맑은 고딕", size=11, bold=True)
        c.alignment = center
        c.border = border
        c.fill = PatternFill("solid", fgColor="EDF3F8")
    ws.row_dimensions[header_row].height = 25

    # 데이터
    for i, t in enumerate(trips):
        r = first_data + i
        vals = [t.get("emp_no", ""), t.get("traveler", ""), t.get("dept", ""),
                t.get("rank", ""), t.get("start_dt_text", ""), t.get("end_dt_text", ""),
                t.get("purpose", ""), t.get("destination", ""),
                t.get("company_car", "X"), t.get("total", 0), t.get("daily", 0),
                t.get("meal", 0), t.get("transport", 0), t.get("travel_time_text", ""),
                round(float(t.get("distance_km", 0)), 1), t.get("fuel_cost", 0),
                t.get("toll_fee", 0), t.get("parking_fee", 0)]
        for j, v in enumerate(vals):
            c = ws.cell(r, j + 2, v)
            c.font = Font(name="맑은 고딕", size=11)
            c.border = border
            c.alignment = left if j in (6, 7) else center
            if j in (9, 10, 11, 12, 15, 16, 17):               # 금액 열
                c.number_format = MONEY_FMT
            elif j == 14:
                c.number_format = "#,##0.0"
        ws.row_dimensions[r].height = 24

    # 총합계
    ws.cell(total_row, 10, "총합계").font = Font(name="맑은 고딕", size=11, bold=True)
    ws.cell(total_row, 10).alignment = center
    tc = ws.cell(total_row, 11, f"=SUM(K{first_data}:K{last_data})")
    tc.font = Font(name="맑은 고딕", size=11, bold=True)
    tc.number_format = MONEY_FMT
    tc.alignment = right
    ws.cell(total_row, 10).border = border
    tc.border = border

    # ---- 계좌 섹션 ----
    title_row = total_row + 1
    ws.merge_cells(f"B{title_row}:G{title_row}")
    ws[f"B{title_row}"] = "여비 지급 개인별 계좌번호"
    ws[f"B{title_row}"].font = Font(name="맑은 고딕", size=20)
    ws[f"B{title_row}"].alignment = center
    for col in "BCDEFG":
        ws[f"{col}{title_row}"].border = border
    ws.row_dimensions[title_row].height = 22

    acct_head = title_row + 3
    for col, label in (("B", "성 명"), ("C", "은 행"), ("F", "예금주"), ("G", "합 계")):
        c = ws[f"{col}{acct_head}"]
        c.value = label
        c.font = Font(name="맑은 고딕", size=11, bold=True)
        c.alignment = center
        c.border = border
    ws.merge_cells(f"D{acct_head}:E{acct_head}")
    ws[f"D{acct_head}"] = "계좌번호"
    ws[f"D{acct_head}"].font = Font(name="맑은 고딕", size=11, bold=True)
    ws[f"D{acct_head}"].alignment = center
    for col in "DE":
        ws[f"{col}{acct_head}"].border = border

    # 성명별 그룹 (등장 순서 유지)
    order: list[str] = []
    groups: dict[str, list[int]] = {}
    for i, t in enumerate(trips):
        name = t.get("traveler", "")
        if name not in groups:
            groups[name] = []
            order.append(name)
        groups[name].append(first_data + i)

    ar = acct_head + 1
    for name in order:
        rows = groups[name]
        ws[f"B{ar}"] = name
        ws[f"C{ar}"] = trips[rows[0] - first_data].get("bank", "")
        ws.merge_cells(f"D{ar}:E{ar}")
        ws[f"D{ar}"] = trips[rows[0] - first_data].get("account", "")
        ws[f"F{ar}"] = trips[rows[0] - first_data].get("holder", name)
        rng = ",".join(f"K{r}" for r in rows)
        ws[f"G{ar}"] = f"=SUM({rng})"
        for col in "BCDEFG":
            cc = ws[f"{col}{ar}"]
            cc.font = Font(name="맑은 고딕", size=11)
            cc.border = border
            cc.alignment = center if col in "BCFG" else left
        ws[f"G{ar}"].number_format = MONEY_FMT
        ar += 1

    ws[f"G{ar}"] = f"=SUM(G{acct_head + 1}:G{ar - 1})"
    ws[f"G{ar}"].font = Font(name="맑은 고딕", size=11, bold=True)
    ws[f"G{ar}"].number_format = MONEY_FMT
    ws[f"G{ar}"].border = border
    ws[f"G{ar}"].alignment = right

    ws.page_setup.orientation = "landscape"
    ws.page_setup.paperSize = ws.PAPERSIZE_A4
    ws.page_setup.fitToWidth = 1          # 인쇄 시 가로 1페이지에 맞춤
    ws.page_setup.fitToHeight = 0
    ws.sheet_properties.pageSetUpPr.fitToPage = True
    ws.print_title_rows = f"{header_row}:{header_row}"
    ws.freeze_panes = f"B{first_data}"

    # ---- 시트2: 산출근거 ----
    ws2 = wb.create_sheet("산출근거")
    heads2 = ["성명", "출장구분", "경로", "차종/유종", "기준연비", "적용유가(원)", "주행거리(km)",
              "연료비(원)", "통행료(원)", "주차료(원)", "식비(원)", "일비(원)", "총액(원)",
              "출장시간", "경로엔진"]
    for i, h in enumerate(heads2):
        c = ws2.cell(1, i + 1, h)
        c.font = Font(name="맑은 고딕", size=10, bold=True)
        c.alignment = center
        c.border = border
        c.fill = PatternFill("solid", fgColor="EDF3F8")
    for r, t in enumerate(trips, start=2):
        row = [t.get("traveler", ""), t.get("scope", ""), t.get("route_text", ""),
               t.get("fuel_name", ""), f"{t.get('efficiency', 0):,.2f} {t.get('eff_unit', '')}",
               round(float(t.get("unit_price", 0)), 2), round(float(t.get("distance_km", 0)), 1),
               t.get("fuel_cost", 0), t.get("toll_fee", 0), t.get("parking_fee", 0),
               t.get("meal", 0), t.get("daily", 0), t.get("total", 0),
               t.get("travel_time_text", ""), t.get("route_provider", "")]
        for j, v in enumerate(row):
            c = ws2.cell(r, j + 1, v)
            c.font = Font(name="맑은 고딕", size=10)
            c.border = border
            c.alignment = left if j in (2,) else center
            if j in (7, 8, 9, 10, 11, 12):
                c.number_format = MONEY_FMT
    for i, w in enumerate([10, 14, 44, 12, 12, 12, 12, 12, 12, 12, 11, 11, 12, 11, 16], start=1):
        ws2.column_dimensions[get_column_letter(i)].width = w
    ws2.freeze_panes = "A2"

    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


# ==============================================================================
# 10. PDF 증빙 보고서
# ==============================================================================

_FONT_PAIRS = [
    (ASSETS_DIR / "NanumGothic.ttf", ASSETS_DIR / "NanumGothicBold.ttf"),
    (BASE_DIR / "NanumGothic.ttf", BASE_DIR / "NanumGothicBold.ttf"),
    (Path("/usr/share/fonts/truetype/nanum/NanumGothic.ttf"),
     Path("/usr/share/fonts/truetype/nanum/NanumGothicBold.ttf")),
    (Path("/usr/share/fonts/truetype/nanum/NanumBarunGothic.ttf"),
     Path("/usr/share/fonts/truetype/nanum/NanumBarunGothicBold.ttf")),
]


def _register_fonts() -> tuple[str, str]:
    from reportlab.pdfbase import pdfmetrics
    from reportlab.pdfbase.ttfonts import TTFont
    for reg, bold in _FONT_PAIRS:
        if reg.exists():
            pdfmetrics.registerFont(TTFont("KFont", str(reg)))
            pdfmetrics.registerFont(TTFont("KFontB", str(bold if bold.exists() else reg)))
            return "KFont", "KFontB"
    return "Helvetica", "Helvetica-Bold"


def build_pdf_report(ctx: dict, map_png: Optional[bytes] = None) -> bytes:
    from reportlab.lib import colors
    from reportlab.lib.enums import TA_CENTER, TA_LEFT
    from reportlab.lib.pagesizes import A4
    from reportlab.lib.styles import ParagraphStyle
    from reportlab.lib.units import mm
    from reportlab.platypus import (Image, KeepTogether, Paragraph, SimpleDocTemplate,
                                    Spacer, Table, TableStyle)

    font, font_b = _register_fonts()
    NAVY = colors.HexColor("#1F3B57")
    GREY = colors.HexColor("#5A6672")
    LINE = colors.HexColor("#C9D2DA")
    BAND = colors.HexColor("#EEF3F7")
    ACCENT = colors.HexColor("#0F5D8C")

    st_title = ParagraphStyle("t", fontName=font_b, fontSize=19, leading=24,
                              alignment=TA_CENTER, textColor=NAVY, spaceAfter=2)
    st_sub = ParagraphStyle("s", fontName=font, fontSize=9, leading=13,
                            alignment=TA_CENTER, textColor=GREY)
    st_h = ParagraphStyle("h", fontName=font_b, fontSize=11.5, leading=15,
                          textColor=ACCENT, spaceBefore=10, spaceAfter=5)
    st_p = ParagraphStyle("p", fontName=font, fontSize=9.3, leading=14, alignment=TA_LEFT)
    st_pb = ParagraphStyle("pb", fontName=font_b, fontSize=9.3, leading=14)
    st_small = ParagraphStyle("sm", fontName=font, fontSize=7.8, leading=11, textColor=GREY)

    def kv(rows, widths=(38 * mm, 138 * mm)) -> Table:
        data = [[Paragraph(k, st_pb), Paragraph(str(v), st_p)] for k, v in rows]
        t = Table(data, colWidths=list(widths))
        t.setStyle(TableStyle([
            ("BACKGROUND", (0, 0), (0, -1), BAND),
            ("GRID", (0, 0), (-1, -1), 0.5, LINE),
            ("VALIGN", (0, 0), (-1, -1), "MIDDLE"),
            ("LEFTPADDING", (0, 0), (-1, -1), 6), ("RIGHTPADDING", (0, 0), (-1, -1), 6),
            ("TOPPADDING", (0, 0), (-1, -1), 4), ("BOTTOMPADDING", (0, 0), (-1, -1), 4)]))
        return t

    story: list = [Paragraph("자가용 출장비 정산 증빙 보고서", st_title),
                   Paragraph("근거: 여비업무 처리 매뉴얼 — 3.운임 나) 자가용 · "
                             "2.여비지급기준 · 4.식비 · 5.일비", st_sub),
                   Spacer(1, 7),
                   Table([[""]], colWidths=[176 * mm], rowHeights=[1.6],
                         style=TableStyle([("BACKGROUND", (0, 0), (-1, -1), NAVY)])),
                   Spacer(1, 8)]

    story.append(Paragraph("1. 출장 개요", st_h))
    story.append(kv([
        ("출장자", f"{ctx['traveler'] or '(미입력)'}"
                  + (f"   ({ctx['dept']})" if ctx.get("dept") else "")
                  + (f"   사번 {ctx['emp_no']}" if ctx.get("emp_no") else "")),
        ("출장구분", ctx["scope"]),
        ("출장기간", f"{ctx['start_dt_text']} ~ {ctx['end_dt_text']}  ({ctx['days']}일간)"),
        ("출장시간", ctx["travel_time_text"]),
        ("출장목적", ctx.get("purpose") or "-"),
        ("출장경로", ctx["route_text"]),
    ]))

    story.append(Paragraph("2. 출장 경로 지도", st_h))
    if map_png:
        from PIL import Image as PILImage
        try:
            with PILImage.open(io.BytesIO(map_png)) as im:
                ratio = im.height / im.width
        except Exception:
            ratio = 0.66
        w = 168 * mm
        h = w * ratio
        if h > 168 * mm:
            h = 168 * mm
            w = h / ratio
        story.append(Image(io.BytesIO(map_png), width=w, height=h))
        story.append(Spacer(1, 3))
        story.append(Paragraph(
            f"※ 실제 주행 경로({ctx.get('route_provider', '')})를 좌표 기반으로 표시. "
            "지도상 경계는 행정구역 경계(시도), 회색 점선은 비교 경로 후보입니다.", st_small))
    else:
        story.append(Paragraph("(경로 지도를 생성하지 못했습니다. 거리·경로는 아래 표를 참조하세요.)", st_p))

    story.append(Paragraph("3. 차량 및 운행 정보", st_h))
    story.append(kv([
        ("차종/유종", ctx["fuel_name"]),
        ("기준 연비·전비", f"{ctx['efficiency']:,.2f} {ctx['eff_unit']}"),
        ("총 주행거리", f"{ctx['distance_km']:,.1f} km   (소요 {ctx['travel_time_text']})"),
        ("적용 유가", f"{ctx['unit_price']:,.2f} {ctx['price_unit']}"),
        ("유가 출처", f"{ctx['price_source']} (출장 시작일 기준)"),
    ]))

    story.append(Paragraph("4. 비용 산정 내역", st_h))
    body: list[list[str]] = []
    if ctx["scope"] == "근무지내 국내출장":
        body.append(["출장비", ctx["within_work_note"] + " (점심시간 포함)", f"{ctx['within_work_cost']:,}"])
        body.append(["통행료", "불가피한 유료도로 통행료 (해당 시)", f"{ctx['toll_fee']:,}"])
    else:
        body.append(["연료비",
                     f"{ctx['distance_km']:,.1f} km ÷ {ctx['efficiency']:,.2f} {ctx['eff_unit']}"
                     f" × {ctx['unit_price']:,.2f} {ctx['price_unit']} ({ctx['rounding_mode']})",
                     f"{ctx['fuel_cost']:,}"])
        body.append(["고속도로 통행료", "통행영수증 실비", f"{ctx['toll_fee']:,}"])
        body.append(["주차료",
                     (f"1일 상한 25,000원 × {ctx['days']}일 적용") if ctx["parking_capped"]
                     else "영수증 실비 (1일 상한 25,000원)", f"{ctx['parking_fee']:,}"])
        if ctx["include_allowances"]:
            body.append(["식비", ctx["meal_note"], f"{ctx['meal_cost']:,}"])
            body.append(["일비", f"1일 25,000원 × {ctx['days']}일", f"{ctx['daily_cost']:,}"])
    body.append(["청구 총액", "합계", f"{ctx['total']:,}"])

    last = len(body) - 1
    data = [[Paragraph(h, st_pb) for h in ["구분", "산출 근거", "금액(원)"]]] + \
           [[Paragraph(c, st_pb if i == last else st_p) for c in row]
            for i, row in enumerate(body)]
    t = Table(data, colWidths=[30 * mm, 111 * mm, 35 * mm])
    t.setStyle(TableStyle([
        ("BACKGROUND", (0, 0), (-1, 0), BAND),
        ("GRID", (0, 0), (-1, -1), 0.5, LINE),
        ("ALIGN", (2, 1), (2, -1), "RIGHT"),
        ("VALIGN", (0, 0), (-1, -1), "MIDDLE"),
        ("BACKGROUND", (0, last + 1), (-1, last + 1), colors.HexColor("#DCEBF6")),
        ("LINEABOVE", (0, last + 1), (-1, last + 1), 1.1, ACCENT),
        ("TOPPADDING", (0, 0), (-1, -1), 5), ("BOTTOMPADDING", (0, 0), (-1, -1), 5)]))
    story.append(t)
    story.append(Spacer(1, 5))
    story.append(Paragraph(
        "※ 주차료는 1일 상한액(일비 25,000원)을 초과할 수 없습니다."
        + ("  본 건은 상한액을 적용하였습니다." if ctx["parking_capped"] else ""), st_small))

    tail = [Paragraph("5. 첨부 증빙서류", st_h)]
    for c in ("[ ] 고속도로 통행영수증" + ("  (해당)" if ctx["toll_fee"] else "  (미해당)"),
              "[ ] 출장지 소재 주유소 결제 신용카드매출전표 (연료비)",
              "[ ] 주차영수증" + ("  (해당)" if ctx["parking_fee"] else "  (미해당)")):
        tail.append(Paragraph(c, st_p))
    tail += [Spacer(1, 3),
             Paragraph("※ 자가용 동승자에게는 연료비·통행료·주차료를 지급하지 않으며, "
                       "2인 이상 동행 출장 시 1대 차량 이용이 원칙입니다.", st_small),
             Paragraph("6. 확인", st_h)]
    sign = Table([[Paragraph("출장자", st_pb), Paragraph("확인자(부서장)", st_pb)],
                  [Paragraph("<br/><br/>(서명 또는 인)", st_p),
                   Paragraph("<br/><br/>(서명 또는 인)", st_p)]],
                 colWidths=[88 * mm, 88 * mm])
    sign.setStyle(TableStyle([("GRID", (0, 0), (-1, -1), 0.5, LINE),
                              ("BACKGROUND", (0, 0), (-1, 0), BAND),
                              ("ALIGN", (0, 0), (-1, -1), "CENTER"),
                              ("TOPPADDING", (0, 0), (-1, -1), 5),
                              ("BOTTOMPADDING", (0, 0), (-1, -1), 5)]))
    tail.append(sign)
    story.append(KeepTogether(tail))
    story.append(Spacer(1, 8))
    story.append(Paragraph(
        f"본 보고서는 자가용 출장비 자동 계산기로 {dt.datetime.now():%Y-%m-%d %H:%M}에 생성되었습니다. "
        "산출 금액은 참고용이며, 최종 지급액은 소속 부서의 증빙서류 확인을 거쳐 확정됩니다.", st_small))

    buf = io.BytesIO()
    SimpleDocTemplate(buf, pagesize=A4, leftMargin=17 * mm, rightMargin=17 * mm,
                      topMargin=15 * mm, bottomMargin=15 * mm,
                      title="자가용 출장비 정산 증빙 보고서",
                      author="출장비 자동 계산기").build(story)
    return buf.getvalue()


# ==============================================================================
# 11. Streamlit UI
# ==============================================================================

def _stretch(widget_func) -> dict:
    try:
        if "width" in inspect.signature(widget_func).parameters:
            return {"width": "stretch"}
    except (TypeError, ValueError):
        pass
    return {"use_container_width": True}


def _init_state() -> None:
    today = dt.date.today()
    for k, v in {"route_options": [], "route_logs": [], "route_choice": 0, "result": None,
                 "map_png": None, "distance_km": 0.0, "toll_fee": 0, "unit_price": 0.0,
                 "dep_date": today, "dep_time": dt.time(9, 0),
                 "arr_date": today, "arr_time": dt.time(18, 0),
                 "time_auto": False, "trips": []}.items():
        st.session_state.setdefault(k, v)


def main() -> None:
    st.set_page_config(page_title="자가용 출장비 자동 계산기", page_icon="🚗", layout="wide")
    _init_state()

    opinet_key = get_secret("OPINET_CERTKEY")
    kakao_key = get_secret("KAKAO_REST_API_KEY")
    vworld_key = get_secret("VWORLD_API_KEY")
    vworld_domain = get_secret("VWORLD_DOMAIN")
    juso_key = get_secret("JUSO_API_KEY")

    st.title("🚗 자가용 출장비 자동 계산기 · 증빙 보고서(PDF·엑셀)")
    st.caption("근거: 여비업무 처리 매뉴얼(2025.09.23.) — 3.운임 나) 자가용 · 2.여비지급기준 · "
               "4.식비 · 5.일비")

    with st.expander("🔑 API 연동 상태 (무료 구성)", expanded=False):
        st.markdown(
            f"""
| 용도 | 사용 API | 상태 |
|---|---|---|
| 실시간 유가 | 오피넷(한국석유공사) | {'✅ 연동' if opinet_key else '⚠️ 키 없음'} |
| 경로·거리·통행료 | 카카오내비 길찾기 | {'✅ 연동' if kakao_key else '⚠️ 키 없음 → OSM 폴백'} |
| 주소→좌표 | VWorld Geocoder 2.0 (무료 40,000건/일) | {'✅ 연동' if vworld_key else '— 미사용'} |
| 주소→좌표 | 도로명주소 API (무료) | {'✅ 연동' if juso_key else '— 미사용'} |
| 주소→좌표 | OSM Nominatim (무료·키 불필요) | ✅ 기본 사용 |

**카카오맵(Local/지오코딩)은 사용하지 않습니다.** 유료 서비스 활성화가 필요 없으며,
주소 변환은 위 **무료 API**로 처리됩니다. VWorld/도로명주소 키를 넣으면 주소 정확도가 더 올라갑니다
([VWorld 인증키](https://www.vworld.kr) · [도로명주소 승인키](https://business.juso.go.kr)).
"""
        )

    # ------------------------------------------------------------ 1) 출장 정보
    st.subheader("1) 출장 정보")
    c1, c2, c3, c4, c5 = st.columns([1.1, 1.1, 0.9, 0.9, 1.1])
    with c1:
        traveler = st.text_input("출장자 성명", placeholder="예: 홍길동")
    with c2:
        dept = st.text_input("부서", placeholder="예: ○○부서")
    with c3:
        emp_no = st.text_input("사번", placeholder="예: 1234567")
    with c4:
        rank = st.text_input("직급", placeholder="예: 기술 5급")
    with c5:
        purpose = st.text_input("사유(용무)", placeholder="예: 협력 워크숍 참석")

    c1, c2, c3, c4 = st.columns(4)
    st.markdown("**출발 일시 / 도착 일시**")
    c1, c2, c3, c4, c5 = st.columns(5)
    with c1:
        dep_date = st.date_input("출발 일자", value=st.session_state["dep_date"])
    with c2:
        dep_time = st.time_input("출발 시간", value=st.session_state["dep_time"],
                                 step=dt.timedelta(minutes=10))
    with c3:
        arr_date = st.date_input("도착 일자", value=st.session_state["arr_date"])
    with c4:
        arr_time = st.time_input("도착 시간", value=st.session_state["arr_time"],
                                 step=dt.timedelta(minutes=10))
    with c5:
        scope = st.selectbox("출장 구분", TRIP_SCOPES, index=0)
    c1, c2, c3 = st.columns([1, 1, 3])
    with c1:
        rounding_mode = st.selectbox("금액 처리 방식", ROUNDING_MODES, index=0)
    with c2:
        auto_arrive = st.checkbox(
            "경로 소요시간으로 도착시간 자동",
            value=False,
            help="체크하면 경로를 계산할 때 도착시간을 경로 소요시간에 맞춰 자동 변경합니다. "
                 "해제(기본)하면 입력한 도착시간을 그대로 사용합니다.",
        )
    with c3:
        st.write("")
        st.caption("도착 시간이 출발 시간보다 빠르면 **다음 날 도착**으로 자동 계산합니다 "
                   "(예: 09:00 출발 → 01:00 도착 = 16시간)")

    dep_dt = dt.datetime.combine(dep_date, dep_time)
    arr_dt = dt.datetime.combine(arr_date, arr_time)
    if arr_dt < dep_dt:                      # 자정 넘김 보정
        arr_dt += dt.timedelta(days=1)
    travel_hours = (arr_dt - dep_dt).total_seconds() / 3600.0
    days = max((arr_dt.date() - dep_dt.date()).days + 1, 1)
    st.caption(f"⏱ 출발 **{dep_dt:%Y.%m.%d %H:%M}** → 도착 **{arr_dt:%Y.%m.%d %H:%M}**  ·  "
               f"출장시간 **{hours_to_text(travel_hours)}**  ·  {days}일간")
    if scope == "근무지내 국내출장":
        amt, note = calc_within_work_allowance(travel_hours)
        st.caption(f"근무지내 기준: {note} → **{amt:,}원** (점심시간 포함 계산)")

    # ------------------------------------------------------------ 2) 경로
    st.subheader("2) 출장 경로 (주소만 입력)")
    c1, c2 = st.columns([1, 1])
    with c1:
        origin = st.text_input("출발지 주소", placeholder="예: 경북 김천시 혁신로 288-7")
        destination = st.text_input("도착지 주소", placeholder="예: 서울역")
    with c2:
        waypoint_raw = st.text_area("경유지 주소 (한 줄에 하나씩, 최대 5개)", height=112,
                                    placeholder="예:\n대전")
    waypoints = [w.strip() for w in waypoint_raw.splitlines() if w.strip()]

    c1, c2, c3 = st.columns([1, 1, 1])
    with c1:
        prio_label = st.selectbox("경로 탐색 기준", list(ROUTE_PRIORITIES), index=0)
    with c2:
        kakao_key_input = st.text_input("카카오 REST API 키 (미입력 시 secrets.toml)",
                                        value="", type="password")
    with c3:
        st.write("")
        calc_route = st.button("📍 경로·거리 자동 계산", type="primary", **_stretch(st.button))
    kakao_key = kakao_key_input.strip() or kakao_key

    if calc_route:
        if not origin.strip() or not destination.strip():
            st.error("출발지와 도착지를 모두 입력해 주세요.")
        else:
            with st.spinner("주소를 좌표로 변환하고 경로 후보를 조회하는 중..."):
                routes, logs = resolve_routes(origin, waypoints, destination, kakao_key,
                                              vworld_key, juso_key,
                                              ROUTE_PRIORITIES[prio_label], vworld_domain)
            st.session_state["route_logs"] = logs
            st.session_state["route_options"] = routes
            st.session_state["route_choice"] = 0
            st.session_state["map_png"] = None

    if st.session_state["route_logs"]:
        with st.expander("경로 계산 로그", expanded=bool(st.session_state["route_options"])):
            for line in st.session_state["route_logs"]:
                st.write(line)

    opts = st.session_state.get("route_options") or []
    if opts:
        st.markdown("**경로 후보 비교** — 확인 후 적용할 경로를 선택하세요")
        cand = pd.DataFrame([{
            "후보": f"{i + 1}", "구분": o.get("priority", ""),
            "거리(km)": round(o["distance_km"], 1),
            "소요시간": hours_to_text(o["duration_min"] / 60.0),
            "통행료(원)": f"{o['toll_fee']:,}"} for i, o in enumerate(opts)])
        st.dataframe(cand, hide_index=True, **_stretch(st.dataframe))
        labels = [f"{i + 1}) {o['distance_km']:.1f}km · {hours_to_text(o['duration_min'] / 60.0)}"
                  f" · 통행료 {o['toll_fee']:,}원" for i, o in enumerate(opts)]
        idx = st.radio("적용할 경로 선택", range(len(opts)),
                       index=min(st.session_state.get("route_choice", 0), len(opts) - 1),
                       format_func=lambda i: labels[i], horizontal=True)
        st.session_state["route_choice"] = idx
        chosen = opts[idx]
        st.session_state["distance_km"] = round(chosen["distance_km"], 1)
        st.session_state["toll_fee"] = int(chosen["toll_fee"])
        if auto_arrive:      # 옵션을 켠 경우에만 도착시간을 경로 소요시간으로 맞춘다
            _dep = dt.datetime.combine(st.session_state["dep_date"], st.session_state["dep_time"])
            _arr = _dep + dt.timedelta(minutes=chosen["duration_min"])
            _new_date, _new_time = _arr.date(), _arr.time().replace(second=0, microsecond=0)
            if (st.session_state["arr_date"] != _new_date
                    or st.session_state["arr_time"] != _new_time):
                st.session_state["arr_date"] = _new_date
                st.session_state["arr_time"] = _new_time
                st.session_state["time_auto"] = True
                st.rerun()
        st.caption(f"🛣️ 경로 소요시간은 약 **{hours_to_text(chosen['duration_min'] / 60)}** 입니다. "
                   "식비·일비는 위에서 입력한 **출발·도착 시간**을 기준으로 계산됩니다.")
        if st.session_state["map_png"] is None:
            with st.spinner("경로 지도를 그리는 중..."):
                st.session_state["map_png"] = render_route_map(
                    chosen, [origin, *waypoints, destination],
                    title="출장 경로 지도 (선택 경로)", routes_all=opts)
        if st.session_state.get("map_png"):
            st.image(st.session_state["map_png"],
                     caption="빨간 실선: 선택 경로 / 회색 점선: 비교 후보", **_stretch(st.image))

    # ------------------------------------------------------ 3) 차종·유종·유가
    st.subheader("3) 차종·유종 및 기준 유가 (오피넷 실시간)")
    c1, c2, c3 = st.columns([1.1, 1, 1])
    with c1:
        fuel_name = st.selectbox("차종/유종 (공단 여비 매뉴얼 연비 기준)", list(FUEL_SPECS))
        spec = FUEL_SPECS[fuel_name]
        st.caption(f"기준 연비·전비 **{spec.efficiency:,.2f} {spec.unit}**"
                   + (f" · {spec.note}" if spec.note else ""))
    with c2:
        sido_options = ["전국"] + list(fetch_opinet_by_sido(opinet_key).keys())
        sido = st.selectbox("유가 지역", sido_options, index=0)
    with c3:
        if st.button("🔄 오피넷 유가 조회", **_stretch(st.button)):
            p = opinet_price(opinet_key, spec.opinet_prodcd, "" if sido == "전국" else sido)
            if p:
                st.session_state["unit_price"] = round(p, 2)
                st.success(f"오피넷 {sido} {spec.name} {p:,.2f} {spec.price_unit} 반영")
            else:
                st.warning(f"{spec.name}은(는) 오피넷에서 조회할 수 없습니다. ({spec.price_source})")

    if not st.session_state["unit_price"]:
        auto_p = opinet_price(opinet_key, spec.opinet_prodcd, "" if sido == "전국" else sido)
        st.session_state["unit_price"] = round(auto_p, 2) if auto_p else spec.fallback_price

    unit_price = st.number_input(f"기준 유가 ({spec.price_unit})", min_value=0.0, step=1.0,
                                 format="%.2f", key="unit_price",
                                 help=f"출처: {spec.price_source} · 출장 시작일 기준 고시가")
    if spec.opinet_prodcd:
        _op = fetch_opinet_all(opinet_key).get(spec.opinet_prodcd)
        if _op:
            st.caption(f"오피넷 기준일 {_op['date']} · {_op['name']} 전국 평균 "
                       f"{_op['price']:,.2f}원 (전일 {_op['diff']}원)")

    # ------------------------------------------------------ 4) 거리·실비
    st.subheader("4) 주행 거리 및 실비")
    c1, c2, c3 = st.columns(3)
    with c1:
        distance_km = st.number_input("총 주행 거리(km)", min_value=0.0, step=1.0,
                                      format="%.1f", key="distance_km")
    with c2:
        toll_fee = st.number_input("고속도로 통행료(원)", min_value=0, step=100, key="toll_fee")
    with c3:
        parking_fee = st.number_input("주차료(원)", min_value=0, step=500, value=0,
                                      help=f"1일 상한 {PARKING_DAILY_CAP:,}원 자동 적용")

    include_allowances = False
    if scope == "근무지외 국내출장":
        include_allowances = st.checkbox("식비·일비도 함께 계산 (규정 제4조)", value=True)

    # 엑셀 계좌정보
    with st.expander("🏦 엑셀 '계좌번호' 칸에 들어갈 정보 (선택)"):
        b1, b2, b3 = st.columns(3)
        with b1:
            bank = st.text_input("은행", placeholder="예: 신한은행")
        with b2:
            account = st.text_input("계좌번호", placeholder="예: 123-456-789012")
        with b3:
            holder = st.text_input("예금주", placeholder="미입력 시 성명과 동일")
        m1, m2 = st.columns([1, 3])
        with m1:
            budget_label = st.text_input("예산과목 라벨", value="예산과목")
        with m2:
            budget_value = st.text_input("예산과목 값",
                                         placeholder="예: 사업운영비-사업비-조사연구비 / 미래형자동차연구개발-연구용역비-조사연구비")
        list_title = st.text_input("엑셀 제목", value=f"{dt.date.today():%Y년 %m월} 국내출장 여비 집행 리스트")

    c1, c2 = st.columns(2)
    with c1:
        calc_now = st.button("🧮 출장비 자동 계산", type="primary", **_stretch(st.button))
    with c2:
        add_trip = st.button("➕ 정산 목록에 추가 (엑셀용)", **_stretch(st.button))

    def _compute() -> Optional[dict]:
        if not traveler.strip():
            st.error("출장자 성명을 입력해 주세요.")
            return None
        if not origin.strip() or not destination.strip():
            st.error("출발지와 도착지를 모두 입력해 주세요.")
            return None
        if travel_hours <= 0:
            st.error("도착 일시는 출발 일시보다 뒤여야 합니다.")
            return None
        if scope == "근무지외 국내출장" and distance_km <= 0:
            st.error("총 주행 거리(km)를 0보다 크게 입력해 주세요.")
            return None
        if scope == "근무지외 국내출장" and unit_price <= 0:
            st.error("기준 유가를 0보다 크게 입력해 주세요.")
            return None

        fuel_cost = (calc_fuel_cost(distance_km, spec.efficiency, unit_price, rounding_mode)
                     if scope == "근무지외 국내출장" else 0)
        parking_applied, parking_capped, cap_total = apply_parking_cap(parking_fee, days)
        toll_applied = int(max(toll_fee, 0))
        transport_total = fuel_cost + toll_applied + parking_applied
        meal_cost, meal_note = calc_meal_allowance(travel_hours, days, MEAL_ALLOWANCE_PER_DAY,
                                                   rounding_mode)
        daily_cost = calc_daily_allowance(days)
        within_cost, within_note = calc_within_work_allowance(travel_hours)
        if scope == "근무지내 국내출장":
            total = within_cost + toll_applied
            excel_daily, excel_meal, excel_transport = 0, 0, toll_applied
        else:
            total = transport_total + ((meal_cost + daily_cost) if include_allowances else 0)
            excel_daily = daily_cost if include_allowances else 0
            excel_meal = meal_cost if include_allowances else 0
            excel_transport = transport_total
        return {
            "traveler": traveler, "dept": dept, "emp_no": emp_no, "rank": rank,
            "purpose": purpose, "trip_date": dep_dt.date(), "end_date": arr_dt.date(),
            "days": days,
            "scope": scope, "travel_hours": travel_hours,
            "travel_time_text": hours_to_text(travel_hours),
            "start_dt_text": f"{dep_dt:%Y.%m.%d %H:%M}", "end_dt_text": f"{arr_dt:%Y.%m.%d %H:%M}",
            "route_text": build_route_text(origin, waypoints, destination),
            "destination": destination,
            "fuel_name": fuel_name, "efficiency": spec.efficiency, "eff_unit": spec.unit,
            "price_unit": spec.price_unit, "price_source": spec.price_source,
            "unit_price": float(unit_price), "distance_km": float(distance_km),
            "fuel_cost": fuel_cost, "toll_fee": toll_applied, "parking_fee": parking_applied,
            "parking_capped": parking_capped, "cap_total": cap_total,
            "transport_total": transport_total, "meal_cost": meal_cost, "meal_note": meal_note,
            "daily_cost": daily_cost, "include_allowances": include_allowances,
            "within_work_cost": within_cost, "within_work_note": within_note,
            "total": total, "rounding_mode": rounding_mode,
            "excel_total": total, "excel_daily": excel_daily, "excel_meal": excel_meal,
            "excel_transport": excel_transport,
            "company_car": "X",
            "bank": bank, "account": account, "holder": holder or traveler,
            "route_provider": (opts[st.session_state["route_choice"]]["provider"] if opts else ""),
            "origin": origin, "waypoints": waypoints,
        }

    if calc_now:
        r = _compute()
        st.session_state["result"] = r
    if add_trip:
        r = _compute()
        if r:
            st.session_state["result"] = r
            st.session_state["trips"].append(r)
            st.success(f"정산 목록에 추가했습니다. (현재 {len(st.session_state['trips'])}건)")

    result = st.session_state.get("result")
    if not result:
        st.info("출장시간 입력 → 주소 입력 → **경로·거리 자동 계산** → **출장비 자동 계산** 순서로 진행하세요.")
        return

    st.divider()
    st.subheader("📊 계산 결과")
    if result["parking_capped"]:
        st.warning(f"주차료 1일 상한액 적용: {result['parking_fee']:,}원 "
                   f"(1일 {PARKING_DAILY_CAP:,}원 × {result['days']}일)")
    if result["scope"] == "근무지내 국내출장":
        k1, k2, k3 = st.columns(3)
        k1.metric("출장비(정액)", f"{result['within_work_cost']:,} 원")
        k2.metric("통행료", f"{result['toll_fee']:,} 원")
        k3.metric("청구 총액", f"{result['total']:,} 원")
    else:
        k1, k2, k3, k4, k5 = st.columns(5)
        k1.metric("연료비", f"{result['fuel_cost']:,} 원")
        k2.metric("통행료", f"{result['toll_fee']:,} 원")
        k3.metric("주차료", f"{result['parking_fee']:,} 원")
        k4.metric("식비+일비",
                  f"{(result['meal_cost'] + result['daily_cost']) if result['include_allowances'] else 0:,} 원")
        k5.metric("청구 총액", f"{result['total']:,} 원")

    rows = [{"구분": "경로", "항목": "출발지", "내역": result["origin"], "금액(원)": "—"},
            {"구분": "경로", "항목": "경유지",
             "내역": ", ".join(result["waypoints"]) if result["waypoints"] else "없음",
             "금액(원)": "—"},
            {"구분": "경로", "항목": "도착지", "내역": result["destination"], "금액(원)": "—"},
            {"구분": "경로", "항목": "총 주행거리", "내역": f"{result['distance_km']:,.1f} km",
             "금액(원)": "—"},
            {"구분": "시간", "항목": "출장시간", "내역": result["travel_time_text"], "금액(원)": "—"},
            {"구분": "시간", "항목": "기간", "내역": f"{result['start_dt_text']} ~ {result['end_dt_text']}",
             "금액(원)": "—"}]
    if result["scope"] == "근무지외 국내출장":
        rows += [
            {"구분": "차량", "항목": "차종/유종", "내역": result["fuel_name"], "금액(원)": "—"},
            {"구분": "차량", "항목": "적용 유가",
             "내역": f"{result['unit_price']:,.2f} {result['price_unit']} ({result['price_source']})",
             "금액(원)": "—"},
            {"구분": "산정", "항목": f"연료비 ({result['rounding_mode']})",
             "내역": f"{result['distance_km']:,.1f} ÷ {result['efficiency']:,.2f} × "
                     f"{result['unit_price']:,.2f}", "금액(원)": f"{result['fuel_cost']:,}"},
            {"구분": "산정", "항목": "고속도로 통행료", "내역": "영수증 실비",
             "금액(원)": f"{result['toll_fee']:,}"},
            {"구분": "산정", "항목": "주차료",
             "내역": "1일 상한 25,000원 적용" if result["parking_capped"] else "영수증 실비",
             "금액(원)": f"{result['parking_fee']:,}"}]
        if result["include_allowances"]:
            rows += [{"구분": "산정", "항목": "식비", "내역": result["meal_note"],
                      "금액(원)": f"{result['meal_cost']:,}"},
                     {"구분": "산정", "항목": "일비", "내역": f"1일 25,000원 × {result['days']}일",
                      "금액(원)": f"{result['daily_cost']:,}"}]
    else:
        rows += [{"구분": "산정", "항목": "근무지내 출장비", "내역": result["within_work_note"],
                  "금액(원)": f"{result['within_work_cost']:,}"},
                 {"구분": "산정", "항목": "통행료(불가피 유료도로)", "내역": "별도 지급 가능",
                  "금액(원)": f"{result['toll_fee']:,}"}]
    rows.append({"구분": "합계", "항목": "청구 총액", "내역": "합계", "금액(원)": f"{result['total']:,}"})
    df = pd.DataFrame(rows)
    st.dataframe(df, hide_index=True, **_stretch(st.dataframe))

    # --------------------------------------------------- PDF / 엑셀 / CSV
    st.subheader("📄 증빙 보고서 (PDF · 엑셀 · CSV)")
    trips = st.session_state.get("trips") or [result]
    meta = {"title": list_title, "budget_label": budget_label, "budget_value": budget_value}
    c1, c2, c3 = st.columns(3)
    try:
        pdf_bytes = build_pdf_report(result, st.session_state.get("map_png"))
        with c1:
            st.download_button("⬇️ 증빙 보고서 PDF", data=pdf_bytes,
                               file_name=f"자가용출장비_증빙보고서_{result['trip_date']:%Y%m%d}_"
                                         f"{(result['traveler'] or '출장자')}.pdf",
                               mime="application/pdf", type="primary",
                               **_stretch(st.download_button))
    except Exception as exc:  # noqa: BLE001
        st.error(f"PDF 생성 오류: `{type(exc).__name__}: {exc}`")
    try:
        xlsx = build_excel(trips, meta)
        with c2:
            st.download_button("⬇️ 출장비 집행 리스트 엑셀", data=xlsx,
                               file_name=f"국내출장_여비집행리스트_{dt.date.today():%Y%m%d}.xlsx",
                               mime="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
                               type="primary", **_stretch(st.download_button))
    except Exception as exc:  # noqa: BLE001
        st.error(f"엑셀 생성 오류: `{type(exc).__name__}: {exc}`")
    with c3:
        st.download_button("⬇️ 정산 내역 CSV", data=df.to_csv(index=False).encode("utf-8-sig"),
                           file_name=f"자가용출장비_정산내역_{result['trip_date']:%Y%m%d}.csv",
                           mime="text/csv", **_stretch(st.download_button))
    st.caption(f"엑셀에는 **정산 목록 {len(trips)}건**이 들어갑니다"
               + (" (➕ 정산 목록에 추가로 여러 건 누적 가능)" if len(trips) == 1 else "")
               + ". 시트 구성: `출장비 집행 리스트`(제목·예산과목·표·총합계·계좌번호) + `산출근거`")

    if st.session_state["trips"]:
        with st.expander(f"📋 정산 목록 ({len(st.session_state['trips'])}건)", expanded=True):
            st.dataframe(pd.DataFrame([{
                "성명": t["traveler"], "사번": t["emp_no"], "기간": t["start_dt_text"],
                "행선지": t["destination"], "출장여비": f"{t['total']:,}",
                "일비": f"{t['excel_daily']:,}", "식비": f"{t['excel_meal']:,}",
                "교통비": f"{t['excel_transport']:,}"} for t in st.session_state["trips"]]),
                hide_index=True, **_stretch(st.dataframe))
            if st.button("🗑 정산 목록 비우기"):
                st.session_state["trips"] = []
                st.rerun()

    st.subheader("🧾 정산 내역 (지출결의서 첨부·참고용)")
    st.text_area("아래 내용을 복사하여 지출결의서에 첨부하세요.",
                 value=build_receipt_text(result), height=440)

    with st.expander("📎 규정 요약 · 증빙서류 체크리스트"):
        st.markdown(
            f"""
**여비 규정 요약**
- **근무지내 국내출장**: 출장시간 **4시간 미만 10,000원 / 4시간 이상 20,000원** (점심시간 포함).
  운임·일비·식비·숙박비 별도 지급 없음. 단, 불가피한 유료도로 통행료는 별도 지급 가능.
- **식비**: 4시간 미만 **1/3** · 4~6시간 **2/3** · 6시간 이상 **전액** (1일 25,000원)
- **일비**: 1일 25,000원 정액
- **자가용 운임**: 연료비(거리÷연비×유가) + 통행료 + 주차료(1일 상한 {PARKING_DAILY_CAP:,}원)

**증빙서류**
- 고속도로 **통행영수증**, 출장지 소재 주유소 **신용카드매출전표**, **주차영수증**
- 자가용 **동승자**에게는 연료비·통행료·주차료를 지급하지 않습니다.
- 2인 이상 동행 출장 시 **1대 차량 이용이 원칙**입니다.
            """
        )


if __name__ == "__main__":
    main()
